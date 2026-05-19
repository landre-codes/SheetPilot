"""Lógica de negócio para exportação de dados — separada da camada GUI."""
import csv
import json
import re
from collections import OrderedDict

_MESES = ["", "jan", "fev", "mar", "abr", "mai", "jun",
          "jul", "ago", "set", "out", "nov", "dez"]

_REGEX_DATA = re.compile(
    r"^(?:\d{2}[/-]\d{2}[/-]\d{4}|\d{2}[/-]\d{4}|\d{4}[/-]\d{2}(?:[/-]\d{2})?|\d{4})$"
)

# ── Queries ─────────────────────────────────────────────────────────────

ARQUIVOS_SQL = """
               SELECT DISTINCT da.id as arq_id, da.nome_arquivo
               FROM fact_dados f
                        JOIN dim_arquivo da ON f.arquivo_id = da.id
               ORDER BY da.nome_arquivo \
               """

SHEETS_POR_ARQUIVO_SQL = """
                         SELECT DISTINCT f.sheet_name
                         FROM fact_dados f
                         WHERE f.arquivo_id = ?
                         ORDER BY f.sheet_name \
                         """

COLUNAS_SQL = """
              SELECT DISTINCT dc.nome_coluna_original
              FROM fact_dados f
                       JOIN dim_coluna dc ON f.coluna_id = dc.id
              WHERE f.arquivo_id = ?
                AND f.sheet_name = ?
              ORDER BY dc.nome_coluna_original \
              """

VALORES_SQL = """
              SELECT f.linha_origem,
                     dc.nome_coluna_original,
                     COALESCE(f.valor_numerico, f.valor_texto) as valor
              FROM fact_dados f
                       JOIN dim_coluna dc ON f.coluna_id = dc.id
              WHERE f.arquivo_id = ?
                AND f.sheet_name = ?
              ORDER BY f.linha_origem, dc.nome_coluna_original \
              """


# ── Utilitários ─────────────────────────────────────────────────────────


def arquivo_id_para_raw(cursor, dim_arq_id: int) -> int | None:
    """Mapeia dim_arquivo.id → arquivos.id para queries em raw_data."""
    row = cursor.execute(
        "SELECT caminho_completo FROM dim_arquivo WHERE id = ?", (dim_arq_id,)
    ).fetchone()
    if not row:
        return None
    try:
        return int(row["caminho_completo"].replace("db://", ""))
    except (ValueError, AttributeError):
        return None


def formatar_data_col(nome: str) -> str:
    """Converte nome de coluna-data para formato legível.
    - dd/mm/aaaa ou aaaa-mm-dd -> preserva a data completa (ex: 15/01/2025)
    - mm/aaaa ou aaaa-mm -> mês/ano abreviado (ex: jan/25)
    - apenas aaaa -> ano (ex: 2025)
    - dd/mm sem ano -> mês/dia
    """
    sep = r"[/-]"
    # dd/mm/aaaa ou dd-mm-aaaa -> preserva data completa
    m = re.match(rf"(\d{{2}}){sep}(\d{{2}}){sep}(\d{{4}})", nome)
    if m:
        return f"{m.group(1)}/{m.group(2)}/{m.group(3)}"
    # aaaa-mm-dd ou aaaa/mm/dd -> converte para dd/mm/aaaa
    m = re.match(rf"(\d{{4}}){sep}(\d{{2}}){sep}(\d{{2}})", nome)
    if m:
        return f"{m.group(3)}/{m.group(2)}/{m.group(1)}"
    # mm/aaaa ou mm-aaaa -> mês/ano abreviado
    m = re.match(rf"(\d{{2}}){sep}(\d{{4}})$", nome)
    if m:
        return f"{_MESES[int(m.group(1))]}/{m.group(2)[2:]}"
    # aaaa-mm ou aaaa/mm -> mês/ano abreviado
    m = re.match(rf"(\d{{4}}){sep}(\d{{2}})$", nome)
    if m:
        return f"{_MESES[int(m.group(2))]}/{m.group(1)[2:]}"
    # dd/mm sem ano -> mês/dia
    m = re.match(rf"(\d{{2}}){sep}(\d{{2}})$", nome)
    if m:
        return f"{m.group(2)}/{m.group(1)}"
    return nome


# ── Construção de tabelas corrigidas ────────────────────────────────────


def build_corrected(valores: list) -> tuple:
    """Constrói formato corrigido a partir dos valores da fact table.
    - Tabelas invertidas (datas como colunas): [entidades, Data, Valor]
    - Tabelas normais: mantém estrutura original (pass-through)."""
    col_set: OrderedDict[str, bool] = OrderedDict()
    for row in valores:
        col_set[row["nome_coluna_original"]] = True

    colunas_entidade: list[str] = []
    colunas_data_raw: list[str] = []
    for col in col_set:
        if _REGEX_DATA.match(col):
            colunas_data_raw.append(col)
        else:
            colunas_entidade.append(col)

    mapa = {(r["linha_origem"], r["nome_coluna_original"]): r["valor"] for r in valores}
    todas_linhas = sorted(set(r["linha_origem"] for r in valores))

    if not colunas_data_raw:
        cabecalho = colunas_entidade
        linhas: list = [cabecalho]
        for ln in todas_linhas:
            row = [str(mapa.get((ln, c), "")) for c in colunas_entidade]
            linhas.append(row)
        return cabecalho, linhas

    cabecalho = colunas_entidade + ["Data", "Valor"]
    linhas = [cabecalho]
    for ln in todas_linhas:
        ent_vals = [str(mapa.get((ln, c), "")) for c in colunas_entidade]
        for col_data in colunas_data_raw:
            dt_fmt = formatar_data_col(col_data)
            val = mapa.get((ln, col_data), "")
            if val is not None and str(val).strip():
                linhas.append(ent_vals + [dt_fmt, str(val)])
    return cabecalho, linhas


def stack_side_by_side(subs: list, colunas_sep: int = 3) -> tuple | None:
    """Empilha múltiplas tabelas lado a lado com N colunas vazias entre elas.
    subs = [(sheet_name, colunas, linhas), ...]"""
    if not subs:
        return None
    all_data = []
    max_data_rows = 0
    for _, _, linhas in subs:
        data = linhas[1:] if len(linhas) > 1 else []
        all_data.append((linhas[0] if linhas else [], data))
        max_data_rows = max(max_data_rows, len(data))

    combined_cols: list[str] = []
    combined_rows: list[list[str]] = [[] for _ in range(max_data_rows)]

    for idx, (header, data) in enumerate(all_data):
        if idx > 0:
            for _ in range(colunas_sep):
                combined_cols.append("")
                for ri in range(max_data_rows):
                    combined_rows[ri].append("")
        combined_cols.extend(header)
        for ri in range(max_data_rows):
            if ri < len(data):
                combined_rows[ri].extend(data[ri])
            else:
                combined_rows[ri].extend([""] * len(header))
    return combined_cols, combined_rows


def build_raw_table(cursor, raw_arquivo_id: int, sheet_name: str) -> tuple:
    """Reconstrói tabela original de raw_data para planilhas normais."""
    rows = cursor.execute("""
                          SELECT rd.row_data
                          FROM raw_data rd
                                   JOIN sheets s ON rd.sheet_id = s.id
                          WHERE rd.arquivo_id = ?
                            AND s.nome_sheet = ?
                          ORDER BY rd.linha_id
                          """, (raw_arquivo_id, sheet_name)).fetchall()
    if not rows:
        return [], []
    first = json.loads(rows[0]["row_data"])
    colunas = list(first.keys())
    data = [colunas]
    for r in rows:
        row_data = json.loads(r["row_data"])
        data.append([str(row_data.get(c, "")) for c in colunas])
    return colunas, data


# ── Escrita de arquivos ─────────────────────────────────────────────────


def _format_xlsx_ws(ws, colunas: list[str], linhas: list):
    """Aplica formatação padronizada + dados a uma planilha openpyxl."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    hf = Font(name="Segoe UI", bold=True, color="FFFFFF", size=11)
    hfill = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")
    ha = Alignment(horizontal="center", vertical="center")
    bdr = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    for ci, col in enumerate(colunas, 1):
        cell = ws.cell(row=1, column=ci, value=col)
        cell.font = hf
        cell.fill = hfill
        cell.alignment = ha
        cell.border = bdr
    ws.freeze_panes = "A2"
    df = Font(name="Segoe UI", size=10)
    af = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
    for ri, row in enumerate(linhas[1:], start=2):
        for ci, val in enumerate(row, 1):
            cell = ws.cell(row=ri, column=ci, value=val if val is not None else "")
            cell.font = df
            cell.border = bdr
            if ri % 2 == 0:
                cell.fill = af
    for ci in range(1, len(colunas) + 1):
        col_val = str(colunas[ci - 1] or "")
        if col_val.strip() == "":
            ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = 1.5
        else:
            ml = len(col_val)
            for ri in range(2, min(len(linhas), 200)):
                v = ws.cell(row=ri, column=ci).value
                if v:
                    ml = max(ml, min(len(str(v)), 40))
            ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = ml + 3


def export_xlsx(path, grupos: list):
    """Exporta para .xlsx. grupos = [(nome_aba, colunas, linhas), ...]"""
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for s_name, colunas, linhas in grupos:
        ws = wb.create_sheet(title=s_name[:31])
        _format_xlsx_ws(ws, colunas, linhas)
    wb.save(str(path))


def export_xlsb(path, grupos: list) -> bool:
    """Exporta para .xlsb via Excel COM. Retorna True se bem-sucedido."""
    try:
        import win32com.client as win32
        excel = win32.gencache.EnsureDispatch("Excel.Application")
        excel.DisplayAlerts = False
        wb = excel.Workbooks.Add()
        while wb.Worksheets.Count > 1:
            wb.Worksheets(2).Delete()
        for idx, (s_name, colunas, linhas) in enumerate(grupos):
            ws = wb.Worksheets(1) if idx == 0 else wb.Worksheets(Add=wb.Worksheets(wb.Worksheets.Count))
            ws.Name = s_name[:31]
            for j, col in enumerate(colunas, 1):
                c = ws.Cells(1, j)
                c.Value = col
                c.Font.Bold = True
                c.Interior.Color = 0x1A237E
                c.Font.Color = 0xFFFFFF
            for i, row in enumerate(linhas[1:], 2):
                for j, val in enumerate(row, 1):
                    ws.Cells(i, j).Value = val if val is not None else ""
            ws.Columns.AutoFit()
        wb.SaveAs(str(path), FileFormat=50)
        wb.Close()
        excel.Quit()
        return True
    except Exception:
        return False


def export_csv(path, colunas: list[str], linhas: list):
    """Exporta para .csv (separador ;)."""
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        for row in linhas:
            w.writerow([str(v) if v is not None else "" for v in row])


def export_parquet(path, colunas: list[str], linhas: list):
    """Exporta para .parquet."""
    import pyarrow as pa
    import pyarrow.parquet as pq
    arrays = {colunas[i]: [] for i in range(len(colunas))}
    for row in linhas[1:]:
        for i, c in enumerate(colunas):
            arrays[c].append(str(row[i]) if row[i] is not None else "")
    pq.write_table(pa.table(arrays), str(path))
