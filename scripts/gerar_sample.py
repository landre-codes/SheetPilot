#!/usr/bin/env python3
"""Gera planilhas de exemplo para testar o pipeline."""
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SAMPLE_DIR = Path(__file__).parent.parent / "data" / "sample"
SAMPLE_DIR.mkdir(parents=True, exist_ok=True)

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def gerar_planilha_normal(nome: str, qtd_linhas: int = 50):
    """Planilha tradicional: linhas = registros, colunas = atributos."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"

    headers = ["ID", "Cliente", "Produto", "Categoria", "Valor", "Data"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    categorias = ["Eletronicos", "Moveis", "Roupas", "Alimentos", "Livros"]
    produtos = ["Prod A", "Prod B", "Prod C", "Prod D", "Prod E"]

    for row in range(2, qtd_linhas + 2):
        ws.cell(row=row, column=1, value=row - 1)
        ws.cell(row=row, column=2, value=f"Cliente {random.randint(1, 20)}")
        ws.cell(row=row, column=3, value=random.choice(produtos))
        ws.cell(row=row, column=4, value=random.choice(categorias))
        ws.cell(row=row, column=5, value=round(random.uniform(10, 5000), 2))
        d = random.randint(1, 28)
        m = random.randint(1, 12)
        ws.cell(row=row, column=6, value=f"{d:02d}/{m:02d}/2025")

    path = SAMPLE_DIR / nome
    wb.save(path)
    print(f"Criado: {path.name} ({qtd_linhas} linhas)")
    return path


def gerar_planilha_invertida(nome: str, qtd_clientes: int = 10, anos: list = None):
    """
    Planilha INVERTIDA (problema tipico):
    Linhas = entidades (clientes)
    Colunas = periodos (datas como cabecalho de coluna)
    """
    if anos is None:
        anos = [2023, 2024, 2025]

    wb = Workbook()
    ws = wb.active
    ws.title = "Vendas"

    headers = ["Cliente", "Regiao"]
    meses = [f"01/{mes:02d}/{ano}" for ano in anos for mes in range(1, 13)]
    headers.extend(meses)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", text_rotation=90)
        cell.number_format = '@'  # força texto para evitar auto-conversão para data

    regioes = ["Norte", "Sul", "Leste", "Oeste", "Centro"]
    for row in range(2, qtd_clientes + 2):
        ws.cell(row=row, column=1, value=f"Cliente {row - 1}")
        ws.cell(row=row, column=2, value=random.choice(regioes))
        for col in range(3, len(headers) + 1):
            ws.cell(row=row, column=col, value=round(random.uniform(100, 50000), 2))

    # Ajusta largura
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 10
    for col in range(3, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 10

    path = SAMPLE_DIR / nome
    wb.save(path)
    print(f"Criado: {path.name} ({qtd_clientes} clientes x {len(headers) - 2} meses)")
    return path


def gerar_planilha_complexa(nome: str):
    """Planilha com multiplas sheets."""
    path_normal = gerar_planilha_normal(f"temp_{nome}", 30)
    path_invertida = gerar_planilha_invertida(f"temp_invertida_{nome}", 5, [2024, 2025])

    # Junta as duas sheets em um arquivo
    wb = Workbook()
    wb.remove(wb.active)

    # Sheet 1: normal
    src_normal = __import__("openpyxl").load_workbook(path_normal)
    ws_src = src_normal.active
    ws_dest = wb.create_sheet("Vendas_Normais")
    for row in ws_src.iter_rows(min_row=1, max_row=ws_src.max_row, max_col=ws_src.max_column, values_only=True):
        ws_dest.append(row)
    src_normal.close()
    path_normal.unlink()

    # Sheet 2: invertida
    src_inv = __import__("openpyxl").load_workbook(path_invertida)
    ws_src = src_inv.active
    ws_dest = wb.create_sheet("Vendas_Periodo")
    for row in ws_src.iter_rows(min_row=1, max_row=ws_src.max_row, max_col=ws_src.max_column, values_only=True):
        ws_dest.append(row)
    src_inv.close()
    path_invertida.unlink()

    path = SAMPLE_DIR / nome
    wb.save(path)
    print(f"Criado: {path.name} (2 sheets combinadas)")


def gerar_planilha_multiplas_tabelas(nome: str):
    """
    Planilha com MÚLTIPLAS TABELAS INVERTIDAS na mesma sheet.
    O verdadeiro problema: cada tabela tem DATAS COMO COLUNAS,
    e elas se repetem verticalmente separadas por linhas em branco.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo_Geral"

    regioes = ["Norte", "Sul", "Leste", "Oeste"]
    meses = [f"01/{m:02d}/2025" for m in range(1, 7)]

    for idx, departamento in enumerate(["Vendas", "Estoque", "Marketing"]):
        if idx > 0:
            ws.append([])
            ws.append([])

        # Cabeçalho: entidade + datas como colunas
        ws.append(["Região"] + meses)

        # Dados: cada linha é uma região, cada coluna é um mês
        for regiao in regioes:
            valores = [round(random.uniform(5000, 80000), 2) for _ in meses]
            ws.append([regiao] + valores)

    path = SAMPLE_DIR / nome
    wb.save(path)
    print(f"Criado: {path.name} (3 tabelas invertidas na mesma sheet, separadas por linhas em branco)")


if __name__ == "__main__":
    print("Gerando planilhas de exemplo...\n")
    gerar_planilha_normal("vendas_normais.xlsx", 100)
    gerar_planilha_invertida("vendas_invertida.xlsx", 15)
    gerar_planilha_complexa("planilha_complexa.xlsx")
    gerar_planilha_multiplas_tabelas("planilha_multiplas_tabelas.xlsx")
    print("\nTodas as planilhas geradas em: data/sample/")
